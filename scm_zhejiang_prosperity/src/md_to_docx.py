"""
将Markdown论文转换为Word文档
"""

from docx import Document
from docx.shared import Pt, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path
import re

def convert_md_to_docx(md_path, docx_path):
    """将Markdown文件转换为Word文档"""
    
    # 读取Markdown文件
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()
    
    # 创建Word文档
    doc = Document()
    
    # 设置页边距
    sections = doc.sections
    for section in sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3.17)
        section.right_margin = Cm(3.17)
    
    # 设置默认字体
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    
    # 解析并转换内容
    lines = content.split('\n')
    i = 0
    
    in_table = False
    table_rows = []
    in_code_block = False
    
    while i < len(lines):
        line = lines[i]
        
        # 跳过数据更新说明（以>开头的引用块）
        if line.startswith('>'):
            i += 1
            continue
        
        # 代码块
        if line.startswith('```'):
            in_code_block = not in_code_block
            i += 1
            continue
        
        if in_code_block:
            # 代码内容用等宽字体
            p = doc.add_paragraph()
            run = p.add_run(line)
            run.font.name = 'Courier New'
            run.font.size = Pt(10)
            i += 1
            continue
        
        # 分隔线
        if line.strip() == '---':
            i += 1
            continue
        
        # 一级标题
        if line.startswith('# '):
            title = line[2:].strip()
            p = doc.add_heading(title, level=0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue
        
        # 二级标题（## ——开头的副标题）
        if line.startswith('## ——'):
            subtitle = line[5:].strip()
            p = doc.add_heading(subtitle, level=1)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue
        
        # 二级标题
        if line.startswith('## '):
            title = line[3:].strip()
            doc.add_heading(title, level=1)
            i += 1
            continue
        
        # 三级标题
        if line.startswith('### '):
            title = line[4:].strip()
            doc.add_heading(title, level=2)
            i += 1
            continue
        
        # 表格
        if line.startswith('|'):
            if not in_table:
                in_table = True
                table_rows = []
            
            # 跳过分隔行
            if '---' in line or ':-' in line or '-:' in line:
                i += 1
                continue
            
            # 解析表格行
            cells = [c.strip() for c in line.split('|')[1:-1]]
            table_rows.append(cells)
            i += 1
            
            # 检查下一行是否还是表格
            if i >= len(lines) or not lines[i].startswith('|'):
                # 表格结束，创建Word表格
                if table_rows:
                    table = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
                    table.style = 'Table Grid'
                    
                    for row_idx, row_data in enumerate(table_rows):
                        for col_idx, cell_data in enumerate(row_data):
                            cell = table.cell(row_idx, col_idx)
                            # 清理Markdown格式
                            clean_text = re.sub(r'\*\*([^*]+)\*\*', r'\1', cell_data)
                            cell.text = clean_text
                    
                    doc.add_paragraph()  # 表格后空行
                
                in_table = False
                table_rows = []
            continue
        
        # 粗体文本行（如摘要、关键词）
        if line.startswith('**') and '**：' in line:
            p = doc.add_paragraph()
            # 解析粗体和普通文本
            match = re.match(r'\*\*([^*]+)\*\*[：:](.*)', line)
            if match:
                bold_text = match.group(1)
                normal_text = match.group(2)
                run = p.add_run(bold_text + '：')
                run.bold = True
                p.add_run(normal_text)
            else:
                p.add_run(line)
            i += 1
            continue
        
        # 图片占位符
        if line.startswith('**[图'):
            p = doc.add_paragraph()
            p.add_run(line.replace('**', ''))
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue
        
        # 普通段落
        if line.strip():
            # 处理Markdown格式
            text = line
            
            # 移除Markdown链接格式
            text = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', text)
            
            # 处理脚注引用 [^n]
            text = re.sub(r'\[\^(\d+)\]', r'[\1]', text)
            
            p = doc.add_paragraph()
            
            # 处理粗体
            parts = re.split(r'(\*\*[^*]+\*\*)', text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    p.add_run(part)
            
            # 首行缩进
            p.paragraph_format.first_line_indent = Cm(0.74)
        
        i += 1
    
    # 保存文档
    doc.save(docx_path)
    print(f"Word文档已保存: {docx_path}")


if __name__ == "__main__":
    PROJECT_ROOT = Path(__file__).parent.parent
    
    # 转换最新版论文
    md_path = PROJECT_ROOT / "paper" / "论文_正式投稿版_2024数据.md"
    docx_path = PROJECT_ROOT / "paper" / "论文_正式投稿版_2024数据.docx"
    
    if md_path.exists():
        convert_md_to_docx(md_path, docx_path)
    else:
        print(f"文件不存在: {md_path}")
